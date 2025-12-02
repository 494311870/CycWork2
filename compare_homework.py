#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
作业对比工具
用于对比两份作业的问题三和问题四答案，找出差异原因
"""

import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

class HomeworkComparator:
    """作业对比类"""
    
    def __init__(self, our_excel, classmate_excel=None):
        """
        初始化
        
        参数:
            our_excel: 我们的Excel文件路径
            classmate_excel: 同桌的Excel文件路径（如果有的话）
        """
        self.our_excel = our_excel
        self.classmate_excel = classmate_excel
        self.doc = Document()
        self.setup_document_style()
        
    def setup_document_style(self):
        """设置文档样式"""
        self.doc.styles['Normal'].font.size = Pt(12)
        
    def analyze_q3_pairs(self, excel_file, label="我们"):
        """
        分析问题三的配对样本
        
        返回:
            配对列表和详细信息
        """
        print(f"\n分析{label}的问题三配对...")
        
        # 读取配对样本
        df_pairs = pd.read_excel(excel_file, sheet_name='Q3数据', header=3, nrows=5)
        
        pairs_info = []
        for i in range(len(df_pairs)):
            row = df_pairs.iloc[i]
            if pd.notna(row['Stkcd']) and pd.notna(row['Stkcd1']):
                pairs_info.append({
                    'index': i + 1,
                    'stkcd': int(row['Stkcd']),
                    'stkcd1': int(row['Stkcd1'])
                })
        
        print(f"{label}的配对数量: {len(pairs_info)}")
        print(f"{label}的配对详情:")
        for pair in pairs_info:
            print(f"  配对{pair['index']}: {pair['stkcd']} ↔ {pair['stkcd1']}")
        
        return pairs_info
    
    def calculate_q3_returns(self, excel_file, pairs_info, label="我们"):
        """
        计算问题三的收益率
        
        参数:
            excel_file: Excel文件路径
            pairs_info: 配对信息列表
            label: 标签
            
        返回:
            每日收益率列表
        """
        print(f"\n计算{label}的问题三收益率...")
        
        # 读取收益率数据
        df_returns = pd.read_excel(excel_file, sheet_name='Q3数据', header=13)
        
        daily_returns = []
        daily_details = []
        
        for day in range(1, 11):
            day_data = df_returns[df_returns['twind'] == day]
            
            pair_returns = []
            valid_pairs = []
            
            for pair in pairs_info:
                inclusion_data = day_data[day_data['Stkcd'] == pair['stkcd']]
                
                if len(inclusion_data) > 0:
                    inclusion_return = inclusion_data.iloc[0]['Dretwd']
                    matched_return = inclusion_data.iloc[0]['Dretwd1']
                    
                    if pd.notna(inclusion_return) and pd.notna(matched_return):
                        pair_return = inclusion_return - matched_return
                        pair_returns.append(pair_return)
                        valid_pairs.append(pair['index'])
            
            if len(pair_returns) > 0:
                portfolio_return = np.mean(pair_returns)
            else:
                portfolio_return = 0
            
            daily_returns.append(portfolio_return * 100)
            daily_details.append({
                'day': day,
                'valid_pairs': len(pair_returns),
                'valid_pair_ids': valid_pairs,
                'return': portfolio_return * 100
            })
        
        # 打印详情
        for detail in daily_details:
            print(f"第{detail['day']}天: 有效配对={detail['valid_pairs']}, "
                  f"配对ID={detail['valid_pair_ids']}, 收益={detail['return']:.4f}%")
        
        cumulative = sum(daily_returns)
        print(f"\n{label}的累计收益: {cumulative:.4f}%")
        
        return daily_returns, daily_details
    
    def analyze_q4_period(self, excel_file, sheet_name, period_name, 
                         pairs_header, pairs_nrows, returns_header, label="我们"):
        """
        分析问题四的某个时期
        
        返回:
            配对信息和收益率
        """
        print(f"\n分析{label}的{period_name}...")
        
        # 读取配对样本
        df_pairs = pd.read_excel(excel_file, sheet_name=sheet_name, 
                                header=pairs_header, nrows=pairs_nrows)
        df_returns = pd.read_excel(excel_file, sheet_name=sheet_name, 
                                  header=returns_header)
        
        # 获取配对信息
        pairs_info = []
        seen_stocks = set()
        
        for i in range(len(df_pairs)):
            row = df_pairs.iloc[i]
            if pd.notna(row['Stkcd']) and pd.notna(row['Stkcd1']):
                stkcd = int(row['Stkcd'])
                if stkcd not in seen_stocks:
                    pairs_info.append({
                        'index': len(pairs_info) + 1,
                        'stkcd': stkcd,
                        'stkcd1': int(row['Stkcd1'])
                    })
                    seen_stocks.add(stkcd)
                if len(pairs_info) >= 5:  # 限制为5对
                    break
        
        print(f"{label}的{period_name}配对数: {len(pairs_info)}")
        
        # 计算10天收益
        daily_returns = []
        for day in range(1, 11):
            day_data = df_returns[df_returns['twind'] == day]
            
            pair_returns = []
            for pair in pairs_info:
                inclusion_data = day_data[day_data['Stkcd'] == pair['stkcd']]
                
                if len(inclusion_data) > 0:
                    inclusion_return = inclusion_data.iloc[0]['Dretwd']
                    matched_return = inclusion_data.iloc[0]['Dretwd1']
                    
                    if pd.notna(inclusion_return) and pd.notna(matched_return):
                        pair_return = inclusion_return - matched_return
                        pair_returns.append(pair_return)
            
            if len(pair_returns) > 0:
                portfolio_return = np.mean(pair_returns)
            else:
                portfolio_return = 0
            
            daily_returns.append(portfolio_return * 100)
        
        cumulative = sum(daily_returns)
        avg = np.mean(daily_returns)
        std = np.std(daily_returns)
        
        print(f"{label}的{period_name}结果:")
        print(f"  累计收益: {cumulative:.4f}%")
        print(f"  平均日收益: {avg:.4f}%")
        print(f"  收益标准差: {std:.4f}%")
        
        return {
            'pairs': pairs_info,
            'daily_returns': daily_returns,
            'cumulative': cumulative,
            'avg': avg,
            'std': std
        }
    
    def compare_q3(self):
        """对比问题三"""
        self.doc.add_heading('问题三对比分析', level=1)
        
        # 分析我们的答案
        our_pairs = self.analyze_q3_pairs(self.our_excel, "我们")
        our_returns, our_details = self.calculate_q3_returns(self.our_excel, our_pairs, "我们")
        
        # 添加我们的结果到文档
        self.doc.add_heading('我们的答案', level=2)
        
        self.doc.add_paragraph(f'配对数量: {len(our_pairs)}对')
        self.doc.add_paragraph('\n配对详情:')
        for pair in our_pairs:
            self.doc.add_paragraph(f'  配对{pair["index"]}: {pair["stkcd"]} ↔ {pair["stkcd1"]}')
        
        self.doc.add_paragraph('\n每日收益率:')
        for i, ret in enumerate(our_returns, 1):
            self.doc.add_paragraph(f'  第{i}天: {ret:.4f}%')
        
        cumulative_our = sum(our_returns)
        self.doc.add_paragraph(f'\n累计收益: {cumulative_our:.4f}%')
        
        # 如果有同桌的答案
        if self.classmate_excel:
            self.doc.add_heading('同桌的答案', level=2)
            
            try:
                classmate_pairs = self.analyze_q3_pairs(self.classmate_excel, "同桌")
                classmate_returns, classmate_details = self.calculate_q3_returns(
                    self.classmate_excel, classmate_pairs, "同桌")
                
                self.doc.add_paragraph(f'配对数量: {len(classmate_pairs)}对')
                self.doc.add_paragraph('\n配对详情:')
                for pair in classmate_pairs:
                    self.doc.add_paragraph(f'  配对{pair["index"]}: {pair["stkcd"]} ↔ {pair["stkcd1"]}')
                
                self.doc.add_paragraph('\n每日收益率:')
                for i, ret in enumerate(classmate_returns, 1):
                    self.doc.add_paragraph(f'  第{i}天: {ret:.4f}%')
                
                cumulative_classmate = sum(classmate_returns)
                self.doc.add_paragraph(f'\n累计收益: {cumulative_classmate:.4f}%')
                
                # 对比分析
                self.doc.add_heading('差异分析', level=2)
                
                # 配对对比
                if len(our_pairs) != len(classmate_pairs):
                    self.doc.add_paragraph(
                        f'⚠️ 配对数量不同: 我们{len(our_pairs)}对 vs 同桌{len(classmate_pairs)}对',
                        style='Intense Quote')
                    self.doc.add_paragraph('这是导致收益率不同的主要原因！')
                else:
                    # 检查配对是否相同
                    pairs_match = True
                    for i, (our_pair, classmate_pair) in enumerate(zip(our_pairs, classmate_pairs)):
                        if (our_pair['stkcd'] != classmate_pair['stkcd'] or 
                            our_pair['stkcd1'] != classmate_pair['stkcd1']):
                            pairs_match = False
                            self.doc.add_paragraph(
                                f'⚠️ 配对{i+1}不同: '
                                f'我们({our_pair["stkcd"]}↔{our_pair["stkcd1"]}) vs '
                                f'同桌({classmate_pair["stkcd"]}↔{classmate_pair["stkcd1"]})')
                    
                    if pairs_match:
                        self.doc.add_paragraph('✅ 配对完全相同')
                
                # 收益率对比
                diff = cumulative_our - cumulative_classmate
                self.doc.add_paragraph(f'\n累计收益差异: {diff:.4f}% (我们 - 同桌)')
                
                if abs(diff) < 0.01:
                    self.doc.add_paragraph('✅ 收益率基本一致（差异 < 0.01%）')
                else:
                    self.doc.add_paragraph(f'⚠️ 收益率存在差异: {abs(diff):.4f}%')
                
            except Exception as e:
                self.doc.add_paragraph(f'❌ 无法读取同桌的数据: {str(e)}')
        else:
            self.doc.add_heading('待对比', level=2)
            self.doc.add_paragraph('请提供同桌的Excel文件以进行对比分析')
    
    def compare_q4(self):
        """对比问题四"""
        self.doc.add_heading('问题四对比分析', level=1)
        
        periods = [
            ('Q4_1数据', '2022年6月（牛市）', 2, 12, 21),
            ('Q4_2数据', '2023年6月（非牛熊市）', 2, 8, 16),
            ('Q4_3数据', '2024年6月（熊市）', 2, 7, 15)
        ]
        
        # 分析我们的答案
        self.doc.add_heading('我们的答案', level=2)
        
        our_results = []
        for sheet_name, period_name, pairs_header, pairs_nrows, returns_header in periods:
            result = self.analyze_q4_period(
                self.our_excel, sheet_name, period_name,
                pairs_header, pairs_nrows, returns_header, "我们")
            our_results.append((period_name, result))
            
            self.doc.add_paragraph(f'\n{period_name}:')
            self.doc.add_paragraph(f'  配对数: {len(result["pairs"])}对')
            self.doc.add_paragraph(f'  累计收益: {result["cumulative"]:.4f}%')
            self.doc.add_paragraph(f'  平均日收益: {result["avg"]:.4f}%')
            self.doc.add_paragraph(f'  收益标准差: {result["std"]:.4f}%')
        
        # 如果有同桌的答案
        if self.classmate_excel:
            self.doc.add_heading('同桌的答案', level=2)
            
            try:
                classmate_results = []
                for sheet_name, period_name, pairs_header, pairs_nrows, returns_header in periods:
                    result = self.analyze_q4_period(
                        self.classmate_excel, sheet_name, period_name,
                        pairs_header, pairs_nrows, returns_header, "同桌")
                    classmate_results.append((period_name, result))
                    
                    self.doc.add_paragraph(f'\n{period_name}:')
                    self.doc.add_paragraph(f'  配对数: {len(result["pairs"])}对')
                    self.doc.add_paragraph(f'  累计收益: {result["cumulative"]:.4f}%')
                    self.doc.add_paragraph(f'  平均日收益: {result["avg"]:.4f}%')
                    self.doc.add_paragraph(f'  收益标准差: {result["std"]:.4f}%')
                
                # 对比分析
                self.doc.add_heading('差异分析', level=2)
                
                for (period_name, our_result), (_, classmate_result) in zip(our_results, classmate_results):
                    self.doc.add_paragraph(f'\n{period_name}对比:')
                    
                    # 配对数对比
                    our_pairs_count = len(our_result['pairs'])
                    classmate_pairs_count = len(classmate_result['pairs'])
                    
                    if our_pairs_count != classmate_pairs_count:
                        self.doc.add_paragraph(
                            f'  ⚠️ 配对数不同: 我们{our_pairs_count}对 vs 同桌{classmate_pairs_count}对',
                            style='Intense Quote')
                    else:
                        self.doc.add_paragraph(f'  ✅ 配对数相同: {our_pairs_count}对')
                    
                    # 收益率对比
                    diff = our_result['cumulative'] - classmate_result['cumulative']
                    self.doc.add_paragraph(
                        f'  累计收益差异: {diff:.4f}% '
                        f'(我们{our_result["cumulative"]:.4f}% vs 同桌{classmate_result["cumulative"]:.4f}%)')
                    
                    if abs(diff) < 0.01:
                        self.doc.add_paragraph('  ✅ 收益率基本一致')
                    else:
                        self.doc.add_paragraph(f'  ⚠️ 收益率存在差异: {abs(diff):.4f}%')
                
            except Exception as e:
                self.doc.add_paragraph(f'❌ 无法读取同桌的数据: {str(e)}')
        else:
            self.doc.add_heading('待对比', level=2)
            self.doc.add_paragraph('请提供同桌的Excel文件以进行对比分析')
    
    def generate_report(self, output_file):
        """生成对比报告"""
        print("\n" + "="*80)
        print("生成对比报告...")
        print("="*80)
        
        # 添加标题
        self.doc.add_heading('作业3问题三和问题四对比分析报告', level=0)
        self.doc.add_paragraph(f'生成时间: {pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")}')
        self.doc.add_paragraph('='*60)
        
        # 对比问题三
        self.compare_q3()
        
        # 对比问题四
        self.compare_q4()
        
        # 保存文档
        self.doc.save(output_file)
        print(f"\n对比报告已保存至: {output_file}")

def main():
    """主函数"""
    print("="*80)
    print("作业对比工具")
    print("="*80)
    
    our_excel = "/home/runner/work/CycWork2/CycWork2/作业3 （学生用）指数纳入效应套利策略的实现.xlsx"
    
    # 同桌的Excel文件（如果有的话）
    classmate_excel = None  # 请替换为同桌的文件路径
    # 例如: classmate_excel = "/path/to/classmate/作业3_指数纳入效应套利策略实现结果(1).xlsx"
    
    output_file = "/home/runner/work/CycWork2/CycWork2/作业对比报告.docx"
    
    # 创建对比对象
    comparator = HomeworkComparator(our_excel, classmate_excel)
    
    # 生成报告
    comparator.generate_report(output_file)
    
    print("\n" + "="*80)
    print("对比完成！")
    print("="*80)
    print(f"\n报告文件: {output_file}")
    
    if classmate_excel is None:
        print("\n⚠️ 注意: 未提供同桌的Excel文件")
        print("如需完整对比，请:")
        print("1. 将同桌的Excel文件放到仓库中")
        print("2. 修改脚本中的 classmate_excel 变量指向该文件")
        print("3. 重新运行此脚本")

if __name__ == "__main__":
    main()
